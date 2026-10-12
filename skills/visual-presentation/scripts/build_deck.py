#!/usr/bin/env python3
"""Build a plain, light 16:9 .pptx deck from a JSON deck spec.

Usage:
    python build_deck.py deck.json out.pptx    # validate, then build
    python build_deck.py deck.json --check     # validate only (no python-pptx needed)

The spec format is described in references/build.md. Building needs python-pptx;
this script never installs anything. Warnings are estimates - always look at the
rendered slides as well (references/qa.md).
"""

import argparse
import json
import math
import sys
from pathlib import Path

# Theme: white ground, near-black text, one accent. All pairs below meet WCAG AA on white.
INK = "1F2937"     # body and headline text, 14.7:1
MUTED = "4B5563"   # secondary text, 7.6:1
LINE = "6B7280"    # arrows, outlines, neutral bars, 4.8:1
PANEL = "F3F4F6"   # light panel fill behind INK text
PAPER = "FFFFFF"
DEFAULT_ACCENT = "0F766E"
DEFAULT_FONT = "Calibri"

SW, SH, M = 13.333, 7.5, 0.6        # slide width/height and side margin, inches
TOP = 1.8                           # top of the visual area
SPLIT_TYPES = {"image", "chart", "number"}

LABELS = {
    "de": {"placeholder": "Bildplatzhalter", "prompt": "Bild-Prompt"},
    "en": {"placeholder": "Image placeholder", "prompt": "Image prompt"},
}

REQUIRED = {
    "title": [], "statement": [], "closing": [],
    "process": ["steps"], "comparison": ["left", "right"], "timeline": ["events"],
    "chart": ["chart"], "number": ["number", "label"], "image": ["image"], "cards": ["cards"],
}


# ---------------------------------------------------------------- validation

def luminance(hex6):
    def channel(c):
        c = c / 255
        return c / 12.92 if c <= 0.03928 else ((c + 0.055) / 1.055) ** 2.4
    r, g, b = (int(hex6[i:i + 2], 16) for i in (0, 2, 4))
    return 0.2126 * channel(r) + 0.7152 * channel(g) + 0.0722 * channel(b)


def contrast(a, b):
    hi, lo = sorted((luminance(a), luminance(b)), reverse=True)
    return (hi + 0.05) / (lo + 0.05)


def fits(text, width_in, height_in, size_pt):
    """Rough estimate: average glyph ~0.5 em wide, line height 1.2."""
    per_line = max(1, int(width_in * 72 / (size_pt * 0.5)))
    lines = sum(max(1, math.ceil(len(p) / per_line)) for p in text.split("\n"))
    return lines * size_pt * 1.2 / 72 <= height_in


def validate(spec):
    errors, warnings = [], []
    accent = spec.get("accent", DEFAULT_ACCENT).lstrip("#")
    if len(accent) != 6:
        errors.append(f"accent '{accent}' must be a 6-digit hex color")
    elif contrast(accent, PAPER) < 4.5:
        warnings.append(f"accent {accent} has {contrast(accent, PAPER):.1f}:1 on white; "
                        "needs 4.5:1 for text (WCAG AA)")
    slides = spec.get("slides") or []
    if not slides:
        errors.append("spec has no slides")
    if len(slides) > 12:
        warnings.append(f"{len(slides)} slides - default is 8-12; confirm the user asked for more")

    for n, s in enumerate(slides, 1):
        where = f"slide {n}"
        kind = s.get("type")
        if kind not in REQUIRED:
            errors.append(f"{where}: unknown type '{kind}'")
            continue
        for field in ["headline"] + REQUIRED[kind]:
            if not s.get(field):
                errors.append(f"{where}: '{field}' is required for type '{kind}'")
        head = s.get("headline", "")
        if head and kind not in ("title", "statement") and len(head.split()) < 4:
            warnings.append(f"{where}: headline '{head}' looks like a topic label - "
                            "write the full statement")
        if len(head) > 90:
            warnings.append(f"{where}: headline has {len(head)} characters (aim for <= 90)")
        lines = s.get("lines", [])
        if len(lines) > 3:
            warnings.append(f"{where}: {len(lines)} text lines - max 3, move details to notes")
        for line in lines:
            if len(line) > 90:
                warnings.append(f"{where}: line '{line[:30]}...' is long (aim for <= 90 chars)")
        if kind != "title" and not s.get("notes"):
            warnings.append(f"{where}: no speaker notes")

        if kind == "process" and not 2 <= len(s.get("steps", [])) <= 5:
            warnings.append(f"{where}: process needs 2-5 steps")
        if kind == "timeline" and not 2 <= len(s.get("events", [])) <= 6:
            warnings.append(f"{where}: timeline needs 2-6 events")
        if kind == "cards" and not 2 <= len(s.get("cards", [])) <= 4:
            warnings.append(f"{where}: cards needs 2-4 cards")
        if kind == "closing" and lines:
            warnings.append(f"{where}: closing ignores 'lines' - put next steps in 'steps'")
        if kind == "closing" and len(s.get("steps", [])) > 3:
            warnings.append(f"{where}: closing should list at most 3 next steps")
        if kind == "number" and len(str(s.get("number", ""))) > 6:
            warnings.append(f"{where}: hero number longer than 6 characters")
        if kind == "comparison":
            for side in ("left", "right"):
                part = s.get(side) or {}
                if not part.get("label"):
                    errors.append(f"{where}: {side}.label is required")
                if len(part.get("items", [])) > 4:
                    warnings.append(f"{where}: {side} has more than 4 items")
        if kind == "chart":
            c = s.get("chart") or {}
            if c.get("kind", "column") not in ("column", "bar", "line"):
                errors.append(f"{where}: chart.kind must be column, bar or line")
            cats, vals = c.get("categories", []), c.get("values", [])
            if not cats or len(cats) != len(vals):
                errors.append(f"{where}: chart needs categories and values of equal length")
            if not all(isinstance(v, (int, float)) for v in vals):
                errors.append(f"{where}: chart values must be numbers")
            if not c.get("source"):
                errors.append(f"{where}: chart.source is required - say where the numbers "
                              "come from; never chart invented data")
        image = s.get("image")
        if image is not None:
            if not image.get("alt"):
                errors.append(f"{where}: image.alt (alt text) is required")
            path = image.get("path")
            if path and not Path(path).is_file():
                errors.append(f"{where}: image file not found: {path}")
            if not path and not image.get("prompt"):
                errors.append(f"{where}: image needs a 'path' or a 'prompt' for the placeholder")
    return errors, warnings


# ---------------------------------------------------------------- building

def build(spec, out_path):
    try:
        from pptx import Presentation
        from pptx.chart.data import CategoryChartData
        from pptx.dml.color import RGBColor
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.enum.dml import MSO_LINE_DASH_STYLE
        from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
        from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
        from pptx.oxml.ns import qn
        from pptx.parts.image import Image
        from pptx.util import Emu, Inches, Pt
        from lxml import etree
    except ImportError:
        sys.exit("python-pptx is not available. Do not install it without asking the user; "
                 "offer the options in references/build.md instead.")

    accent = spec.get("accent", DEFAULT_ACCENT).lstrip("#")
    font = spec.get("font", DEFAULT_FONT)
    lang = spec.get("lang", "en-US")
    words = LABELS.get(lang.split("-")[0], LABELS["en"])
    warnings = []

    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)   # 13.333 x 7.5 in
    props = prs.core_properties
    props.title = spec.get("title", "")
    props.language = lang
    props.author = props.last_modified_by = spec.get("author", "")
    props.comments = props.subject = props.keywords = ""
    layout = prs.slide_layouts[5]   # "Title Only": every slide gets a real title

    def rgb(hex6):
        return RGBColor.from_string(hex6)

    def style_frame(tf, paras, size, color=INK, bold=False, align=PP_ALIGN.LEFT,
                    anchor=MSO_ANCHOR.TOP, spacing=0.35):
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.vertical_anchor = anchor
        tf.margin_left = tf.margin_right = Inches(0.05)
        tf.margin_top = tf.margin_bottom = Inches(0.03)
        for i, text in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.space_after = Pt(size * spacing) if i < len(paras) - 1 else Pt(0)
            run = p.add_run()
            run.text = text
            run.font.name, run.font.size, run.font.bold = font, Pt(size), bold
            run.font.color.rgb = rgb(color)
            run._r.get_or_add_rPr().set("lang", lang)

    def text(slide, x, y, w, h, paras, size, where="", **kw):
        if isinstance(paras, str):
            paras = [paras]
        if not fits("\n".join(paras), w - 0.1, h - 0.06, size):
            warnings.append(f"{where}: text may overflow its box: '{paras[0][:40]}'")
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        style_frame(box.text_frame, paras, size, **kw)
        return box

    def shape(slide, kind, x, y, w, h, fill=None, line=None, line_w=1.0, dash=False):
        sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        if fill:
            sh.fill.solid()
            sh.fill.fore_color.rgb = rgb(fill)
        else:
            sh.fill.background()
        if line:
            sh.line.color.rgb = rgb(line)
            sh.line.width = Pt(line_w)
            if dash:
                sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        return sh

    def label_shape(sh, paras, size, color=INK, bold=False, align=PP_ALIGN.CENTER,
                    anchor=MSO_ANCHOR.MIDDLE, where=""):
        paras = paras if isinstance(paras, list) else [paras]
        if where and not fits("\n".join(paras), sh.width / 914400 - 0.3,
                              sh.height / 914400 - 0.1, size):
            warnings.append(f"{where}: label may overflow its shape: '{paras[0][:40]}'")
        style_frame(sh.text_frame, paras, size,
                    color=color, bold=bold, align=align, anchor=anchor)
        sh.text_frame.margin_left = sh.text_frame.margin_right = Inches(0.12)

    def arrow(slide, x1, y1, x2, y2, color=LINE, width=2.0):
        c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                       Inches(x2), Inches(y2))
        c.line.color.rgb = rgb(color)
        c.line.width = Pt(width)
        tail = etree.SubElement(c.line._get_or_add_ln(), qn("a:tailEnd"))
        tail.set("type", "triangle")
        return c

    def alt(sh, description):
        cnv = sh._element.xpath("./*[1]/p:cNvPr")
        if cnv:
            cnv[0].set("descr", description)

    def headline(slide, s, size=32, y=0.45, h=1.25, w=SW - 2 * M, x=M,
                 anchor=MSO_ANCHOR.TOP):
        title = slide.shapes.title
        title.left, title.top, title.width, title.height = (Inches(x), Inches(y),
                                                            Inches(w), Inches(h))
        title.text_frame.clear()
        style_frame(title.text_frame, [s["headline"]], size, bold=True, anchor=anchor)
        if not fits(s["headline"], w - 0.1, h, size):
            warnings.append(f"{s['_where']}: headline may not fit in two lines - shorten it")

    # ---- visuals: each draws into the box (x, y, w, h)

    def draw_process(slide, s, x, y, w, h):
        steps = [st if isinstance(st, dict) else {"label": st} for st in s["steps"]]
        n, gap = len(steps), 0.7
        bw = (w - (n - 1) * gap) / n
        bh = min(1.5, h * 0.45)
        has_detail = any(st.get("detail") for st in steps)
        by = y + (h - bh - (1.1 if has_detail else 0)) / 2 + 0.2
        for i, st in enumerate(steps):
            bx = x + i * (bw + gap)
            key = s.get("highlight") == i
            box = shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, bx, by, bw, bh,
                        fill=accent if key else PANEL, line=None if key else LINE)
            box.adjustments[0] = 0.12
            label_shape(box, st["label"], 20, color=PAPER if key else INK, bold=True,
                        where=s["_where"])
            if st.get("detail"):
                text(slide, bx, by + bh + 0.15, bw, 1.0, st["detail"], 16, s["_where"],
                     color=MUTED, align=PP_ALIGN.CENTER)
            if i < n - 1:
                arrow(slide, bx + bw + 0.08, by + bh / 2, bx + bw + gap - 0.08, by + bh / 2)
                if st.get("arrow"):
                    text(slide, bx + bw + gap / 2 - 0.8, by - 0.5, 1.6, 0.4, st["arrow"], 14,
                         s["_where"], color=MUTED, align=PP_ALIGN.CENTER)

    def draw_comparison(slide, s, x, y, w, h):
        pw = (w - 0.5) / 2
        for i, side in enumerate(("left", "right")):
            part = s[side]
            px = x + i * (pw + 0.5)
            key = side == s.get("highlight", "right")
            head = shape(slide, MSO_SHAPE.RECTANGLE, px, y, pw, 0.7,
                         fill=accent if key else PANEL)
            label_shape(head, part["label"], 22, color=PAPER if key else INK, bold=True,
                        align=PP_ALIGN.LEFT, where=s["_where"])
            body = shape(slide, MSO_SHAPE.RECTANGLE, px, y + 0.7, pw, h - 0.7,
                         line=accent if key else LINE, line_w=2.0 if key else 1.0)
            label_shape(body, part.get("items", []), 20, align=PP_ALIGN.LEFT,
                        anchor=MSO_ANCHOR.TOP)
            body.text_frame.margin_top = Inches(0.2)
            if not fits("\n".join(part.get("items", [])), pw - 0.3, h - 1.0, 20):
                warnings.append(f"{s['_where']}: {side} items may overflow")

    def draw_timeline(slide, s, x, y, w, h):
        events = s["events"]
        n = len(events)
        ly = y + h * 0.42
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(ly),
                                          Inches(x + w), Inches(ly))
        line.line.color.rgb, line.line.width = rgb(LINE), Pt(2.5)
        slot = w / n
        for i, ev in enumerate(events):
            cx = x + slot * (i + 0.5)
            key = s.get("highlight") == i
            shape(slide, MSO_SHAPE.OVAL, cx - 0.16, ly - 0.16, 0.32, 0.32,
                  fill=accent if key else LINE)
            text(slide, cx - slot / 2 + 0.1, ly - 0.95, slot - 0.2, 0.6, ev["when"], 20,
                 s["_where"], color=accent if key else INK, bold=True, align=PP_ALIGN.CENTER,
                 anchor=MSO_ANCHOR.BOTTOM)
            text(slide, cx - slot / 2 + 0.1, ly + 0.35, slot - 0.2, 1.6, ev["label"], 16,
                 s["_where"], align=PP_ALIGN.CENTER)

    def draw_chart(slide, s, x, y, w, h):
        c = s["chart"]
        data = CategoryChartData()
        data.categories = c["categories"]
        data.add_series(c.get("series", s["headline"][:40]), c["values"])
        kind = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED, "bar": XL_CHART_TYPE.BAR_CLUSTERED,
                "line": XL_CHART_TYPE.LINE_MARKERS}[c.get("kind", "column")]
        frame = slide.shapes.add_chart(kind, Inches(x), Inches(y), Inches(w),
                                       Inches(h - 0.45), data)
        chart = frame.chart
        chart.has_legend = False
        chart.font.size, chart.font.name = Pt(16), font
        chart.font.color.rgb = rgb(INK)
        chart.value_axis.has_major_gridlines = False
        chart.value_axis.visible = False
        plot = chart.plots[0]
        plot.has_data_labels = True
        labels = plot.data_labels
        labels.font.size, labels.font.bold = Pt(16), True
        if c.get("number_format"):
            labels.number_format, labels.number_format_is_linked = c["number_format"], False
        series = plot.series[0]
        hi = c.get("highlight")
        if c.get("kind") == "line":
            series.format.line.color.rgb = rgb(accent)
            series.format.line.width = Pt(3)
        else:
            plot.gap_width = 60
            for i in range(len(c["values"])):
                point = series.points[i]
                point.format.fill.solid()
                point.format.fill.fore_color.rgb = rgb(accent if hi in (None, i) else LINE)
        alt(frame, c.get("alt") or (s["headline"] + ": " + ", ".join(
            f"{k} {v}" for k, v in zip(c["categories"], c["values"]))))
        text(slide, x, y + h - 0.4, w, 0.4, c["source"], 12, s["_where"], color=MUTED)

    def draw_number(slide, s, x, y, w, h):
        text(slide, x, y + 0.3, w, 1.9, str(s["number"]), 96, s["_where"], color=accent,
             bold=True, anchor=MSO_ANCHOR.BOTTOM)
        text(slide, x, y + 2.3, w, 1.4, s["label"], 24, s["_where"])

    def draw_image(slide, img, x, y, w, h, where):
        path = img.get("path")
        if path:
            px_w, px_h = Image.from_file(path).size
            scale = min(w / px_w, h / px_h)
            iw, ih = px_w * scale, px_h * scale
            pic = slide.shapes.add_picture(path, Inches(x + (w - iw) / 2),
                                           Inches(y + (h - ih) / 2), Inches(iw), Inches(ih))
            alt(pic, img["alt"])
            return
        box = shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill="F9FAFB", line=LINE,
                    line_w=1.5, dash=True)
        label_shape(box, [words["placeholder"], img["prompt"]], 14, color=MUTED,
                    align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE)
        box.text_frame.paragraphs[0].runs[0].font.bold = True
        alt(box, words["placeholder"] + ": " + img["alt"])

    def draw_cards(slide, s, x, y, w, h):
        cards = s["cards"]
        n, gap = len(cards), 0.4
        cw = (w - (n - 1) * gap) / n
        ch = min(h, 3.2)
        for i, card in enumerate(cards):
            key = s.get("highlight") == i
            box = shape(slide, MSO_SHAPE.RECTANGLE, x + i * (cw + gap), y, cw, ch,
                        fill=accent if key else PANEL)
            label_shape(box, [card["title"]] + ([card["text"]] if card.get("text") else []),
                        18, color=PAPER if key else INK, align=PP_ALIGN.LEFT,
                        anchor=MSO_ANCHOR.TOP, where=s["_where"])
            box.text_frame.margin_top = Inches(0.25)
            first = box.text_frame.paragraphs[0].runs[0].font
            first.bold, first.size = True, Pt(22)

    def draw_closing(slide, s, x, y, w, h):
        for i, step in enumerate(s.get("steps", [])):
            sy = y + 0.2 + i * 1.15
            dot = shape(slide, MSO_SHAPE.OVAL, x, sy, 0.7, 0.7, fill=accent)
            label_shape(dot, str(i + 1), 20, color=PAPER, bold=True)
            dot.text_frame.margin_left = dot.text_frame.margin_right = Inches(0)
            text(slide, x + 1.0, sy, w - 1.0, 0.75, step, 24, s["_where"],
                 anchor=MSO_ANCHOR.MIDDLE)

    draw = {"process": draw_process, "comparison": draw_comparison,
            "timeline": draw_timeline, "chart": draw_chart, "number": draw_number,
            "cards": draw_cards, "closing": draw_closing}

    # ---- slides

    for n, s in enumerate(spec["slides"], 1):
        s["_where"] = f"slide {n}"
        kind = s["type"]
        slide = prs.slides.add_slide(layout)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(PAPER)
        lines = s.get("lines", [])
        notes = s.get("notes", [])
        notes = [notes] if isinstance(notes, str) else list(notes)

        if kind == "title":
            img = s.get("image")
            tw = 6.3 if img else SW - 2 * M
            headline(slide, s, size=44, y=1.6, h=2.4, w=tw, anchor=MSO_ANCHOR.BOTTOM)
            if s.get("subtitle"):
                text(slide, M, 4.2, tw, 1.1, s["subtitle"], 22, s["_where"], color=MUTED)
            if s.get("meta"):
                text(slide, M, 6.4, tw, 0.5, s["meta"], 16, s["_where"], color=MUTED)
            if img:
                draw_image(slide, img, 7.4, 0.9, SW - M - 7.4, 5.7, s["_where"])
        elif kind == "statement":
            headline(slide, s, size=40, y=1.9, h=2.4, anchor=MSO_ANCHOR.BOTTOM)
            if lines:
                text(slide, M, 4.5, SW - 2 * M, 1.6, lines, 22, s["_where"], color=MUTED)
        else:
            headline(slide, s)
            split = kind in SPLIT_TYPES and lines
            if split:
                vx, vy, vw, vh = M, TOP, 7.6, 5.0
            elif kind in SPLIT_TYPES or kind == "closing":
                vx, vy, vw, vh = M, TOP, SW - 2 * M, 5.0
            else:
                vx, vy, vw, vh = M, TOP, SW - 2 * M, 3.8 if lines else 5.0
            if kind == "image":
                draw_image(slide, s["image"], vx, vy, vw, vh, s["_where"])
            else:
                draw[kind](slide, s, vx, vy, vw, vh)
            if lines and split:
                text(slide, M + 8.0, TOP + 0.2, SW - 2 * M - 8.0, 4.6, lines, 20, s["_where"])
            elif lines and kind != "closing":
                text(slide, M, 5.8, SW - 2 * M, 1.1, lines, 20, s["_where"])
            text(slide, SW - M - 1.0, 6.95, 1.0, 0.35, str(n), 12, s["_where"],
                 color=MUTED, align=PP_ALIGN.RIGHT)

        image = s.get("image")
        if image and not image.get("path"):
            notes.append(f"{words['prompt']}: {image['prompt']}")
        if notes:
            tf = slide.notes_slide.notes_text_frame
            tf.text = "\n".join(notes)
            for p in tf.paragraphs:
                for r in p.runs:
                    r._r.get_or_add_rPr().set("lang", lang)

    prs.save(out_path)
    return warnings


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("spec", help="deck spec (JSON)")
    ap.add_argument("out", nargs="?", help="output .pptx path")
    ap.add_argument("--check", action="store_true", help="validate the spec only")
    args = ap.parse_args()

    spec = json.loads(Path(args.spec).read_text(encoding="utf-8"))
    base = Path(args.spec).resolve().parent
    for s in spec.get("slides", []):          # image paths are relative to the spec file
        img = s.get("image") or {}
        if img.get("path") and not Path(img["path"]).is_absolute():
            img["path"] = str(base / img["path"])

    errors, warnings = validate(spec)
    if not args.check and not errors:
        if not args.out:
            ap.error("output path required unless --check is given")
        if Path(args.out).exists():
            sys.exit(f"{args.out} exists - choose a new name or ask before overwriting")
        warnings += build(spec, args.out)
    for w in warnings:
        print(f"WARNING {w}", file=sys.stderr)
    for e in errors:
        print(f"ERROR {e}", file=sys.stderr)
    if errors:
        sys.exit(2)
    print("spec OK" if args.check else f"wrote {args.out} ({len(spec['slides'])} slides, "
          f"{len(warnings)} warnings)")


if __name__ == "__main__":
    main()
